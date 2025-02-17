import smtplib
import ssl
import time
import hashlib
import openpyxl
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from bs4 import BeautifulSoup
from datetime import datetime

# SMTP configuration
port = 587  # If using STARTTLS, port 587 is typical. (SMTP_SSL normally uses 465.)
username = 'Documents@web-notification.info'
password = 'CX$4567@vajra1$123'
server = '208.91.198.96'

# Load the HTML email template
with open('email.html', 'r', encoding='utf-8') as file:
    soup = BeautifulSoup(file, features="html.parser")
email_template = soup.prettify()

# Define a default URL to replace the {{.URL}} placeholder
default_url = "http://example.com/gift"  # Update this URL as needed

# Load the Excel workbook containing recipient details
workbook = openpyxl.load_workbook('emails.xlsx')
sheet = workbook.active

# Iterate over each row (skipping the header row)
for row in sheet.iter_rows(min_row=2, values_only=True):
    # Unpack the row data: FirstName, LastName, Email, Position
    first_name, last_name, email, position = row
    
    # Skip if email is empty
    if not email:
        continue
    
    # Compute a tracker value (MD5 hash of the email)
    tracker = hashlib.md5(email.encode()).hexdigest()
    
    # Replace placeholders in the email template with actual values
    body = email_template
    body = body.replace('{{.FirstName}}', first_name)
    body = body.replace('{{.LastName}}', last_name)
    body = body.replace('{{.Position}}', position)
    body = body.replace('{{.URL}}', default_url)
    body = body.replace('{{.Tracker}}', tracker)
    
    # Log the sending event with the current date and time
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"Sending email to: {email} at {now} | Tracker: {tracker}")
    
    # Create the email message
    message = MIMEMultipart('alternative')
    message['Subject'] = "Thank You for Your Continued Support! 🎁"  # Update the subject as needed
    message['From'] = username
    message['To'] = email
    message.attach(MIMEText(body, 'html'))
    
    # Send the email using an SSL connection
    try:
        context = ssl.create_default_context()
        # Note: Typically, port 587 is used with SMTP and STARTTLS. If your server supports SSL on 587, this is fine.
        with smtplib.SMTP_SSL(server, port, context=context) as server_conn:
            server_conn.login(username, password)
            server_conn.sendmail(username, email, message.as_string())
    except Exception as e:
        print(f"Error sending email to {email}: {e}")
    
    # Wait for 3 seconds before sending the next email
    time.sleep(3)
